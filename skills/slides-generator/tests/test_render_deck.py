"""Offline contract and OOXML tests. Scratch artifacts stay inside this test folder."""

from copy import deepcopy
import hashlib
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess
import sys
import unittest
from unittest.mock import patch
import uuid
import zipfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "render_deck.py"
SPEC = importlib.util.spec_from_file_location("render_deck", SCRIPT)
renderer = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(renderer)


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        self.outline, self.raw = renderer.load_json(ROOT / "assets" / "template-outline.json")
        self.content, _ = renderer.load_json(ROOT / "assets" / "template-deck.json")
        self.work = ROOT / "tests" / (".run-" + uuid.uuid4().hex)
        self.work.mkdir()
        self.addCleanup(shutil.rmtree, self.work)
        self.approval = {
            "status": "approved",
            "outline_sha256": hashlib.sha256(self.raw).hexdigest(),
            "outline_version": self.outline["version"],
            "user_confirmation": "Да, утверждаю эту структуру. (Test fixture only.)",
        }

    def write(self, name, data):
        path = self.work / name
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return path

    def cli(self, *arguments):
        return subprocess.run(
            [sys.executable, str(SCRIPT), *map(str, arguments)],
            capture_output=True, text=True, encoding="utf-8",
        )

    def inputs(self):
        outline = self.work / "outline.json"
        outline.write_bytes(self.raw)
        return outline, self.write("approval.json", self.approval), self.write("content.json", self.content)

    def assert_invalid(self, mutation, message):
        mutation()
        with self.assertRaisesRegex(renderer.ValidationError, message):
            renderer.render(self.outline, self.content, self.work / "invalid")
        self.assertFalse((self.work / "invalid").exists())

    def test_all_layouts_native_rendering_and_notes(self):
        paths = renderer.render(self.outline, self.content, self.work / "deck", placeholder=True)
        self.assertEqual([p.name for p in paths], ["template.pptx", "template-speech.md", "template-sources.md"])
        presentation = Presentation(paths[0])
        self.assertEqual(len(presentation.slides), 7)
        self.assertEqual(presentation.slide_width, Inches(13.333))
        self.assertEqual(presentation.slide_height, Inches(7.5))
        self.assertEqual(presentation.core_properties.language, "ru")
        self.assertIn("PLACEHOLDER", presentation.core_properties.keywords)
        for slide, content, planned in zip(presentation.slides, self.content["slides"], self.outline["slides"]):
            text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            self.assertIn(content["title"], text)
            self.assertIn("Sources / Источники:", text)
            self.assertFalse(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes))
            notes = slide.notes_slide.notes_text_frame.text
            self.assertIn(content["speech"], notes)
            self.assertIn(content["transition"], notes)
            self.assertIn(f"{planned['seconds']} seconds", notes)
            self.assertIn(renderer.PLACEHOLDER, notes)
            self.assertIn(self.outline["audience"], notes)
        table = next(shape.table for shape in presentation.slides[4].shapes if shape.has_table)
        self.assertEqual(table.cell(1, 1).text, "Условная задача А")
        table.cell(1, 1).text = "Проверка редактирования"
        chart = next(shape.chart for shape in presentation.slides[5].shapes if shape.has_chart)
        self.assertEqual(chart.series[1].values, (-4.0, 14.0, 20.0))
        self.assertEqual(chart.series[0].name, "Пример 1")
        self.assertEqual([category.label for category in chart.plots[0].categories], ["А", "Б", "В"])
        self.assertIsNotNone(chart.part.chart_workbook.xlsx_part)
        updated = BytesIO()
        presentation.save(updated)
        reopened = Presentation(BytesIO(updated.getvalue()))
        self.assertEqual(next(shape.table for shape in reopened.slides[4].shapes if shape.has_table).cell(1, 1).text,
                         "Проверка редактирования")
        for path in paths[1:]:
            self.assertIn(renderer.PLACEHOLDER, path.read_text(encoding="utf-8"))
        self.assertIn(self.content["slides"][0]["speech"], paths[1].read_text(encoding="utf-8"))
        with zipfile.ZipFile(paths[0]) as archive:
            self.assertTrue(any(name.startswith("ppt/embeddings/") for name in archive.namelist()))
            self.assertEqual(len([name for name in archive.namelist()
                                  if name.startswith("ppt/charts/chart") and name.endswith(".xml")]), 1)
            self.assertFalse(any(name.startswith("ppt/media/") for name in archive.namelist()))

    def assert_geometry(self, data):
        presentation = Presentation(BytesIO(data))
        for slide in presentation.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, presentation.slide_width)
                self.assertLessEqual(shape.top + shape.height, presentation.slide_height)
                if shape.has_text_frame and shape.text:
                    frame = shape.text_frame
                    required = 0
                    for paragraph in frame.paragraphs:
                        self.assertIsNotNone(paragraph.font.size)
                        required += paragraph.font.size.pt * 1.15
                        self.assertEqual(paragraph.font.name, "Arial")
                        self.assertGreaterEqual(paragraph.font.size.pt, 12)
                    available = (shape.height - frame.margin_top - frame.margin_bottom) / 12700
                    self.assertLessEqual(required, available + 0.1, shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            required = sum(p.font.size.pt * 1.15 for p in cell.text_frame.paragraphs)
                            available = (row.height - cell.margin_top - cell.margin_bottom) / 12700
                            self.assertLessEqual(required, available + 0.1)
                            self.assertTrue(all(p.font.size.pt >= 12 for p in cell.text_frame.paragraphs))

    def test_geometry_and_readable_fonts(self):
        self.assert_geometry(renderer.build_presentation(self.outline, self.content, placeholder=True))

    def test_chart_category_axis_schema_order_and_singleton_elements(self):
        presentation = Presentation(BytesIO(renderer.build_presentation(self.outline, self.content)))
        chart = next(s.chart for s in presentation.slides[5].shapes if s.has_chart)
        axis = chart.category_axis._element
        self.assertEqual(len(axis.xpath("./c:noMultiLvlLbl")), 1)
        self.assertEqual(len(axis.xpath("./c:tickLblSkip")), 1)
        self.assertEqual(axis.xpath("./c:tickLblSkip")[0].get("val"), "1")
        sequence = (
            "axId", "scaling", "delete", "axPos", "majorGridlines", "minorGridlines",
            "title", "numFmt", "majorTickMark", "minorTickMark", "tickLblPos",
            "spPr", "txPr", "crossAx", "crosses", "crossesAt", "auto", "lblAlgn",
            "lblOffset", "tickLblSkip", "tickMarkSkip", "noMultiLvlLbl", "extLst",
        )
        indices = [sequence.index(element.tag.rsplit("}", 1)[-1]) for element in axis]
        self.assertEqual(indices, sorted(indices))
        self.assertFalse(chart.legend.include_in_layout)

    def test_long_audience_metadata_and_core_title_limit(self):
        self.outline["audience"] = "Я" * 500
        data = renderer.build_presentation(self.outline, self.content)
        presentation = Presentation(BytesIO(data))
        self.assertIn(self.outline["audience"], presentation.slides[0].notes_slide.notes_text_frame.text)
        self.outline["title"] = "x" * 256
        with self.assertRaisesRegex(renderer.ValidationError, "outline.title"):
            renderer.validate_outline(self.outline)

    def test_maximum_layout_sizes_and_negative_chart(self):
        self.content["slides"][1]["body"]["items"] = ["Образец тезиса\nВторая строка"] * 5
        comparison = self.content["slides"][2]["body"]
        comparison["left_items"] = comparison["right_items"] = ["Краткий пример\nВторая строка"] * 4
        table = self.content["slides"][4]["body"]
        table["columns"] = ["Пример\nколонки"] * 5
        table["rows"] = [["Две\nстроки"] * 5 for _ in range(6)]
        chart = self.content["slides"][5]["body"]
        chart["categories"] = ["А", "Б", "В", "Г", "Д", "Е"]
        chart["series"] = [
            {"name": f"Ряд {i}", "values": [-1e9, -0.2, 0, 1e-9, 3, 1e9]} for i in range(3)
        ]
        renderer.validate_content(self.content, self.outline)
        data = renderer.build_presentation(self.outline, self.content)
        self.assert_geometry(data)
        presentation = Presentation(BytesIO(data))
        native_chart = next(s.chart for s in presentation.slides[5].shapes if s.has_chart)
        self.assertEqual(len(native_chart.series), 3)
        self.assertEqual(len(native_chart.series[0].values), 6)
        self.assertLess(native_chart.value_axis.minimum_scale, -1e9)
        self.assertGreater(native_chart.value_axis.maximum_scale, 1e9)

    def test_zero_and_all_negative_chart_ranges(self):
        for values in ([0, 0, 0], [-1, -2, -3]):
            with self.subTest(values=values):
                self.content["slides"][5]["body"]["series"] = [{"name": "Пример", "values": values}]
                data = renderer.build_presentation(self.outline, self.content)
                presentation = Presentation(BytesIO(data))
                chart = next(s.chart for s in presentation.slides[5].shapes if s.has_chart)
                self.assertLess(chart.value_axis.minimum_scale, chart.value_axis.maximum_scale)

    def test_chart_labels_are_literal_strings_in_embedded_workbook(self):
        body = self.content["slides"][5]["body"]
        body["categories"][0] = "=2+2"
        body["series"][0]["name"] = "=3+3"
        renderer.validate_content(self.content, self.outline)
        presentation = Presentation(BytesIO(renderer.build_presentation(self.outline, self.content)))
        chart = next(s.chart for s in presentation.slides[5].shapes if s.has_chart)
        with zipfile.ZipFile(BytesIO(chart.part.chart_workbook.xlsx_part.blob)) as workbook:
            worksheet = workbook.read("xl/worksheets/sheet1.xml").decode("utf-8")
            strings = workbook.read("xl/sharedStrings.xml").decode("utf-8")
            self.assertNotIn("<f", worksheet)
            self.assertNotIn("<hyperlink", worksheet)
            self.assertIn("=2+2", strings)
            self.assertIn("=3+3", strings)
        self.assertEqual(chart.series[0].name, "=3+3")

    def test_hash_command_uses_raw_bytes_without_pptx_dependency(self):
        outline, _, _ = self.inputs()
        outline.write_bytes(self.raw + b"\n")
        result = subprocess.run(
            [sys.executable, "-S", str(SCRIPT), "outline-hash", "--outline", str(outline)],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip(), hashlib.sha256(self.raw + b"\n").hexdigest())
        self.assertNotEqual(result.stdout.strip(), self.approval["outline_sha256"])

    def test_valid_approved_cli_and_overwrite_protection(self):
        outline, approval, content = self.inputs()
        out = self.work / "approved"
        arguments = ["render", "--outline", outline, "--approval", approval, "--content", content, "--output-dir", out]
        result = self.cli(*arguments)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual({p.name for p in out.iterdir()}, {"presentation.pptx", "speech.md", "sources.md"})
        before = {p.name: p.read_bytes() for p in out.iterdir()}
        result = self.cli(*arguments)
        self.assertEqual(result.returncode, 2)
        self.assertIn("already exists", result.stderr)
        self.assertEqual(before, {p.name: p.read_bytes() for p in out.iterdir()})
        result = self.cli(*arguments, "--overwrite")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(len(Presentation(out / "presentation.pptx").slides), 7)

    def test_template_cli_needs_no_approval_but_is_explicitly_placeholder(self):
        result = self.cli("template", "--theme", "light-business", "--output-dir", self.work / "sample")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((self.work / "sample" / "template.pptx").is_file())
        result = self.cli("template", "--theme", "light-business", "--output-dir", self.work / "sample")
        self.assertEqual(result.returncode, 2)
        result = self.cli("template", "--theme", "light-business", "--output-dir", self.work / "sample", "--overwrite")
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_approval_failures(self):
        for key, value in (
            ("status", "draft"),
            ("outline_sha256", "0" * 64),
            ("outline_version", 2),
            ("outline_version", True),
            ("outline_version", 1.0),
            ("user_confirmation", ""),
            ("user_confirmation", "  "),
        ):
            with self.subTest(key=key, value=value):
                approval = {**self.approval, key: value}
                with self.assertRaisesRegex(renderer.ValidationError, "approval." + key):
                    renderer.validate_approval(approval, self.outline, self.raw)
        renderer.validate_approval(self.approval, self.outline, self.raw)
        with self.assertRaisesRegex(renderer.ValidationError, "raw outline bytes"):
            renderer.validate_approval(self.approval, self.outline, self.raw + b" ")

    def test_cli_rejected_approval_leaves_no_outputs(self):
        self.approval["status"] = "pending"
        outline, approval, content = self.inputs()
        result = self.cli("render", "--outline", outline, "--approval", approval,
                          "--content", content, "--output-dir", self.work / "invalid")
        self.assertEqual(result.returncode, 2)
        self.assertIn("approval.status", result.stderr)
        self.assertFalse((self.work / "invalid").exists())

    def test_mutated_and_reordered_content(self):
        for key, value in (("title", "New title"), ("id", "other"), ("layout", "closing")):
            with self.subTest(key=key):
                content = deepcopy(self.content)
                content["slides"][0][key] = value
                with self.assertRaisesRegex(renderer.ValidationError, "exactly match"):
                    renderer.validate_content(content, self.outline)
        self.content["slides"][0], self.content["slides"][1] = self.content["slides"][1], self.content["slides"][0]
        with self.assertRaisesRegex(renderer.ValidationError, "exactly match"):
            renderer.validate_content(self.content, self.outline)

    def test_duplicate_outline_ids_and_slide_count(self):
        self.outline["slides"][1]["id"] = self.outline["slides"][0]["id"]
        with self.assertRaisesRegex(renderer.ValidationError, "duplicate"):
            renderer.validate_outline(self.outline)
        self.content["slides"].pop()
        with self.assertRaisesRegex(renderer.ValidationError, "7–7"):
            renderer.validate_content(self.content, self.outline)

    def test_line_wrapping_not_character_count(self):
        wide, narrow = "W" * 18, "i" * 18
        self.assertEqual(renderer.wrapped(narrow, 2, 18, 1, "test"), narrow)
        with self.assertRaisesRegex(renderer.ValidationError, "word exceeds"):
            renderer.wrapped(wide, 2, 18, 1, "test")
        with self.assertRaisesRegex(renderer.ValidationError, "needs 3 lines"):
            renderer.wrapped("one\ntwo\nthree", 10, 18, 2, "test")

    def test_each_layout_rejects_overflow(self):
        for index, mutation in (
            (0, lambda b: b.update(subtitle="Word " * 100)),
            (1, lambda b: b.update(items=["a"] * 6)),
            (1, lambda b: b.update(items=["Word " * 100])),
            (2, lambda b: b.update(left_items=["x"] * 5)),
            (2, lambda b: b.update(right_title="Long " * 100)),
            (3, lambda b: b["items"][0].update(value="12345678901234567890")),
            (3, lambda b: b["items"][0].update(context="line\n" * 6)),
            (4, lambda b: b["rows"][0].__setitem__(0, "Word " * 100)),
            (4, lambda b: b.update(rows=[["x"] * 3] * 7)),
            (5, lambda b: b.update(categories=["a"] * 7)),
            (5, lambda b: b.update(categories=["a long category label", "Б", "В"])),
            (6, lambda b: b.update(takeaway="Word " * 100)),
            (6, lambda b: b.update(actions=["x"] * 4)),
        ):
            with self.subTest(index=index, mutation=mutation):
                content = deepcopy(self.content)
                mutation(content["slides"][index]["body"])
                with self.assertRaises(renderer.ValidationError):
                    renderer.render(self.outline, content, self.work / "invalid")
                self.assertFalse((self.work / "invalid").exists())

    def test_numeric_validation(self):
        for value in (True, "1", None, float("nan"), float("inf"), -float("inf"), 1e10, 1e-10):
            with self.subTest(value=value):
                self.content["slides"][5]["body"]["series"][0]["values"][0] = value
                with self.assertRaisesRegex(renderer.ValidationError, "values"):
                    renderer.validate_content(self.content, self.outline)
        for value in (True, "60", None, 0, -1, float("inf")):
            with self.subTest(seconds=value):
                self.outline["slides"][0]["seconds"] = value
                with self.assertRaisesRegex(renderer.ValidationError, "seconds"):
                    renderer.validate_outline(self.outline)
        self.outline["slides"][0]["seconds"] = 61
        with self.assertRaisesRegex(renderer.ValidationError, "total exceeds"):
            renderer.validate_outline(self.outline)

    def test_chart_series_lengths_and_names(self):
        chart = self.content["slides"][5]["body"]
        chart["series"][0]["values"] = [1, 2]
        with self.assertRaisesRegex(renderer.ValidationError, "3–3"):
            renderer.validate_content(self.content, self.outline)
        chart["series"][0]["values"] = [1, 2, 3]
        chart["series"][1]["name"] = chart["series"][0]["name"]
        with self.assertRaisesRegex(renderer.ValidationError, "duplicate"):
            renderer.validate_content(self.content, self.outline)

    def source(self):
        return {
            "id": "src01", "title": "Тестовый документ", "publisher": "Тестовая команда",
            "date": "n.d.", "location": "documents/example.pdf", "locator": "стр. 5",
            "access": "local", "accessed_at": "2026-09-06",
        }

    def test_sources_in_footer_notes_and_markdown(self):
        self.content["sources"] = [self.source()]
        self.content["slides"][0]["source_ids"] = ["src01"]
        paths = renderer.render(self.outline, self.content, self.work / "sourced")
        slide = Presentation(paths[0]).slides[0]
        self.assertTrue(any("src01" in shape.text for shape in slide.shapes if shape.has_text_frame))
        notes = slide.notes_slide.notes_text_frame.text
        for field in self.source().values():
            self.assertIn(field, notes)
        self.assertIn("src01", paths[1].read_text(encoding="utf-8"))
        sources = paths[2].read_text(encoding="utf-8")
        self.assertIn("Тестовый документ", sources)
        self.assertIn("Used on slides: s01", sources)

    def test_unresolved_duplicate_and_invalid_source_refs(self):
        self.content["slides"][0]["source_ids"] = ["src01"]
        with self.assertRaisesRegex(renderer.ValidationError, "unresolved"):
            renderer.validate_content(self.content, self.outline)
        self.content["sources"] = [self.source(), self.source()]
        with self.assertRaisesRegex(renderer.ValidationError, "duplicate"):
            renderer.validate_content(self.content, self.outline)
        self.content["sources"] = [self.source()]
        self.content["slides"][0]["source_ids"] = ["src01", "src01"]
        with self.assertRaisesRegex(renderer.ValidationError, "duplicate"):
            renderer.validate_content(self.content, self.outline)

    def test_source_metadata_validation(self):
        for key, value in (("access", "private"), ("date", "2026-02-30"), ("date", "2026-1-01"),
                           ("accessed_at", "n.d."), ("locator", ""), ("location", None)):
            with self.subTest(key=key):
                source = self.source()
                source[key] = value
                self.content["sources"] = [source]
                with self.assertRaises(renderer.ValidationError):
                    renderer.validate_content(self.content, self.outline)

    def test_output_preflight_does_not_partial_write(self):
        out = self.work / "out"
        out.mkdir()
        (out / "speech.md").write_text("keep me", encoding="utf-8")
        with self.assertRaisesRegex(renderer.ValidationError, "already exists"):
            renderer.render(self.outline, self.content, out)
        self.assertEqual([p.name for p in out.iterdir()], ["speech.md"])
        self.assertEqual((out / "speech.md").read_text(encoding="utf-8"), "keep me")
        (out / "speech.md").unlink()
        (out / "sources.md").symlink_to(self.work / "elsewhere")
        with self.assertRaisesRegex(renderer.ValidationError, "symlink"):
            renderer.render(self.outline, self.content, out, overwrite=True)
        self.assertFalse((self.work / "elsewhere").exists())

    def test_failed_serialization_creates_no_outputs(self):
        with patch.object(renderer, "build_presentation", side_effect=ValueError("serialization failed")):
            with self.assertRaisesRegex(ValueError, "serialization failed"):
                renderer.render(self.outline, self.content, self.work / "invalid")
        self.assertFalse((self.work / "invalid").exists())

    def test_missing_dependency_has_actionable_error(self):
        result = subprocess.run(
            [sys.executable, "-S", str(SCRIPT), "template", "--theme", "light-business", "--output-dir", str(self.work / "missing")],
            capture_output=True, text=True,
        )
        self.assertEqual(result.returncode, 2)
        self.assertIn("dependency missing", result.stderr)
        self.assertIn("pip install", result.stderr)
        self.assertIn(str(ROOT / "requirements.txt"), result.stderr)
        self.assertNotIn("Traceback", result.stderr)
        self.assertFalse((self.work / "missing").exists())

    def test_invalid_json_and_missing_files_cli(self):
        outline, approval, content = self.inputs()
        for invalid in ('{"slides": NaN}', '{"slides": [], "slides": []}', "{broken"):
            content.write_text(invalid, encoding="utf-8")
            result = self.cli("render", "--outline", outline, "--approval", approval, "--content", content,
                              "--output-dir", self.work / "invalid")
            self.assertEqual(result.returncode, 2)
            self.assertNotIn("Traceback", result.stderr)
            self.assertFalse((self.work / "invalid").exists())
        result = self.cli("outline-hash", "--outline", self.work / "missing.json")
        self.assertEqual(result.returncode, 2)

    def test_rejects_unsupported_text_and_unknown_fields(self):
        self.content["slides"][0]["speech"] = "Control character: \x00"
        with self.assertRaisesRegex(renderer.ValidationError, "control characters"):
            renderer.validate_content(self.content, self.outline)
        self.content["slides"][0]["speech"] = "\ud800"
        with self.assertRaisesRegex(renderer.ValidationError, "XML"):
            renderer.validate_content(self.content, self.outline)
        self.content["slides"][0]["speech"] = "Test speech."
        self.content["slides"][0]["body"]["ignored"] = "Do not silently ignore data"
        with self.assertRaisesRegex(renderer.ValidationError, "unknown fields"):
            renderer.validate_content(self.content, self.outline)


if __name__ == "__main__":
    unittest.main()
