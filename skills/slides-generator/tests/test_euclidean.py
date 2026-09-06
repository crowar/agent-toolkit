"""Euclidean contract, bundled-font, and native editable OOXML regression tests."""

from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess
import sys
import unittest
from unittest.mock import patch
import uuid
import zipfile

from PIL import Image
from pptx import Presentation

from test_render_deck import ROOT, SCRIPT, renderer


class EuclideanTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.sample_outline, cls.raw = renderer.load_json(ROOT / "assets/euclidean/template-outline.json")
        cls.sample_content, _ = renderer.load_json(ROOT / "assets/euclidean/template-deck.json")
        cls.euclidean = renderer.euclidean()

    def setUp(self):
        self.outline, self.content = deepcopy(self.sample_outline), deepcopy(self.sample_content)
        self.work = ROOT / "tests" / (".run-" + uuid.uuid4().hex)
        self.work.mkdir()
        self.addCleanup(shutil.rmtree, self.work)

    def invalid(self, fragment):
        with self.assertRaisesRegex(renderer.ValidationError, fragment):
            renderer.render(self.outline, self.content, self.work / "invalid")
        self.assertFalse((self.work / "invalid").exists())

    def cli(self, *arguments):
        return subprocess.run([sys.executable, str(SCRIPT), *map(str, arguments)], capture_output=True, text=True)

    def test_all_nineteen_source_layouts_are_native_and_exact(self):
        renderer.validate_outline(self.outline)
        renderer.validate_content(self.content, self.outline)
        payload = renderer.build_presentation(self.outline, self.content, True)
        presentation = Presentation(BytesIO(payload))
        self.assertEqual(len(presentation.slides), 19)
        self.assertEqual(presentation.slide_width, 12192000)
        self.assertEqual(presentation.slide_height, 6858000)
        self.assertIn("DEMO", presentation.core_properties.keywords)
        catalogue = self.euclidean.catalogue()
        for slide, data in zip(presentation.slides, self.content["slides"]):
            layout = catalogue[data["layout"]]
            self.assertEqual(str(slide.background.fill.fore_color.rgb), layout["background"][1:])
            self.assertEqual(len(slide.shapes), len(layout["elements"]))
            notes = slide.notes_slide.notes_text_frame.text
            self.assertIn(data["speech"], notes)
            self.assertIn(data["transition"], notes)
            for ref in data["source_ids"]:
                self.assertIn(f"[{ref}]", notes)
            for shape, element in zip(slide.shapes, layout["elements"]):
                self.assertEqual(shape.name, element["id"])
                self.assertNotIn("Sources /", shape.text if shape.has_text_frame else "")
                self.assertFalse(shape._element.xpath(".//p:pic"))
                self.assertFalse(shape._element.xpath("./p:style"))
                self.assertTrue(shape._element.xpath("./p:spPr/a:effectLst"))
                if element["kind"] == "line":
                    start = self.euclidean.point(element["transform"], 0, 0)
                    end = self.euclidean.point(element["transform"], *element["size"])
                    self.assertAlmostEqual(shape.begin_x, self.euclidean.px(start[0]), delta=1)
                    self.assertAlmostEqual(shape.begin_y, self.euclidean.px(start[1]), delta=1)
                    self.assertAlmostEqual(shape.end_x, self.euclidean.px(end[0]), delta=1)
                    self.assertAlmostEqual(shape.end_y, self.euclidean.px(end[1]), delta=1)
                    for coordinate, limit in ((start[0], 1920), (end[0], 1920), (start[1], 1080), (end[1], 1080)):
                        self.assertGreaterEqual(coordinate, -3)
                        self.assertLessEqual(coordinate, limit + 3)
                else:
                    self.assertEqual(shape.left, self.euclidean.px(element["transform"][4]))
                    self.assertEqual(shape.top, self.euclidean.px(element["transform"][5]))
                    self.assertEqual(shape.width, self.euclidean.px(element["size"][0]))
                    self.assertEqual(shape.height, self.euclidean.px(element["size"][1]))
                if element["kind"] == "text":
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            self.assertEqual(run.font.size.pt, element["font"]["size"] / 2)
                            self.assertEqual(run.font.bold, element["font"]["style"] == "Bold")
                            self.assertTrue(run.font.name.startswith(element["font"]["family"]))
                    if element["slot"] is None:
                        self.assertEqual(shape.text, "“")
                        self.assertEqual(shape.height, 0)
                        self.assertTrue(shape._element.xpath(".//a:alpha"))
                elif element["kind"] == "vector":
                    self.assertTrue(shape._element.xpath(".//a:custGeom/a:pathLst/a:path/a:cubicBezTo"))
                elif element["fill"]:
                    self.assertEqual(str(shape.fill.fore_color.rgb), element["fill"]["color"][1:])
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            self.assertFalse(any(name.startswith("ppt/media/") for name in archive.namelist()))
            slides_xml = "".join(archive.read(name).decode() for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
            self.assertIn('typeface="Source Sans Pro Semibold"', slides_xml)
            self.assertIn('typeface="Source Sans Pro Black"', slides_xml)
            self.assertIn('sz="7000"', slides_xml)
            self.assertNotIn('val="2563EB"', slides_xml)
            self.assertIn("buAutoNum", slides_xml)

    def test_alternative_text_uses_current_content_not_source_metadata(self):
        self.outline["slides"][0]["title"] = self.content["slides"][0]["title"] = "Новая мысль"
        self.content["slides"][0]["body"]["fields"]["caption"] = "Данные"
        renderer.validate_outline(self.outline)
        renderer.validate_content(self.content, self.outline)
        deck = Presentation(BytesIO(renderer.build_presentation(self.outline, self.content)))
        shapes = {shape.name: shape for shape in deck.slides[0].shapes}
        self.assertEqual(shapes["n235"]._element._nvXxPr.cNvPr.get("descr"), "ДАННЫЕ")
        self.assertEqual(shapes["n236"]._element._nvXxPr.cNvPr.get("descr"), "Новая мысль")
        self.assertIsNone(shapes["n237"]._element._nvXxPr.cNvPr.get("descr"))
        for slide in deck.slides:
            for shape in slide.shapes:
                self.assertNotIn('"baselines"', shape._element._nvXxPr.cNvPr.get("descr", ""))

    def test_numbered_list_markers_preserve_source_offsets(self):
        deck = Presentation(BytesIO(renderer.build_presentation(self.outline, self.content)))
        layout = self.euclidean.catalogue()["euclidean-agenda"]
        element = next(e for e in layout["elements"] if e["id"] == "n247")
        shape = next(s for s in deck.slides[2].shapes if s.name == "n247")
        for paragraph, baseline in zip(shape.text_frame.paragraphs, element["baselines"]):
            properties = paragraph._p.pPr
            self.assertEqual(int(properties.get("marL")), self.euclidean.px(baseline["x"]))
            self.assertEqual(int(properties.get("indent")), self.euclidean.px(baseline["marker_x"] - baseline["x"]))
            self.assertAlmostEqual(baseline["marker_x"], 17.6642, places=4)

    def test_default_template_and_persistent_demo(self):
        result = self.cli("template", "--output-dir", self.work / "demo")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(len(Presentation(self.work / "demo/template.pptx").slides), 19)
        persistent = Presentation(ROOT / "assets/template.pptx")
        self.assertEqual(len(persistent.slides), 19)
        self.assertIn("DEMO", persistent.core_properties.keywords)
        for data in self.content["slides"]:
            self.assertIn(data["speech"], (self.work / "demo/template-speech.md").read_text())

    def test_catalog_is_dependency_free_and_reports_slots(self):
        result = subprocess.run([sys.executable, "-S", str(SCRIPT), "catalog", "--theme", "euclidean"], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        catalogue = json.loads(result.stdout)
        self.assertEqual(len(catalogue["layouts"]), 19)
        self.assertEqual(sum(len(layout["image_slots"]) for layout in catalogue["layouts"]), 8)
        self.assertIn("baselines", catalogue["layouts"][0]["elements"][0])

    def test_unknown_theme(self):
        self.outline["theme"] = "other"
        self.invalid("outline.theme")

    def test_design_and_font_manifests_are_explicitly_utf8(self):
        read_text = Path.read_text

        def read_utf8(path, *args, **kwargs):
            self.assertEqual(kwargs.get("encoding"), "utf-8")
            return read_text(path, *args, **kwargs)

        with patch.object(Path, "read_text", autospec=True, side_effect=read_utf8):
            self.euclidean.catalog_document()
            layout = self.euclidean.catalogue()["euclidean-cover"]
            element = next(e for e in layout["elements"] if e["kind"] == "text")
            self.euclidean.font_asset(element)

    def test_cross_theme_layout_rejected(self):
        self.outline["slides"][0]["layout"] = "cover"
        self.invalid("outline.slides\\[0\\].layout")

    def test_omitted_theme_remains_light_business(self):
        del self.outline["theme"]
        self.invalid("outline.slides\\[0\\].layout")

    def test_missing_fields_no_source_fallback(self):
        self.content["slides"][0]["body"]["fields"].pop("caption")
        self.invalid("missing fields: caption")

    def test_unknown_fields(self):
        self.content["slides"][0]["body"]["fields"]["extra"] = "No"
        self.invalid("unknown fields: extra")

    def test_empty_field(self):
        self.content["slides"][0]["body"]["fields"]["caption"] = ""
        self.invalid("must not be empty")

    def test_body_schema(self):
        self.content["slides"][0]["body"]["subtitle"] = "Not a Euclidean slot"
        self.invalid("unknown fields: subtitle")

    def test_actual_title_slot_overflow(self):
        self.outline["slides"][0]["title"] = " ".join(["Too much title content"] * 35)
        self.content["slides"][0]["title"] = self.outline["slides"][0]["title"]
        self.invalid("fixed 70pt")

    def test_field_overflow(self):
        self.content["slides"][0]["body"]["fields"]["caption"] = "A very long date with more words"
        self.invalid("slot is 35px")

    def test_short_replacement_uses_exact_size(self):
        self.outline["slides"][0]["title"] = "A new idea"
        self.content["slides"][0]["title"] = "A new idea"
        self.content["slides"][0]["body"]["fields"]["caption"] = "2026"
        renderer.validate_outline(self.outline)
        renderer.validate_content(self.content, self.outline)
        deck = Presentation(BytesIO(renderer.build_presentation(self.outline, self.content)))
        title = next(shape for shape in deck.slides[0].shapes if shape.name == "n236")
        self.assertEqual(title.text, "A new idea")
        self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.size.pt, 70)

    def test_missing_font_manifest_is_actionable(self):
        with patch.object(self.euclidean, "ROOT", self.work):
            element = next(iter(self.euclidean_catalogue()["euclidean-cover"]["elements"]))
            with self.assertRaisesRegex(renderer.ValidationError, "missing bundled font manifest"):
                self.euclidean.font_asset(element)

    def euclidean_catalogue(self):
        return {layout["id"]: layout for layout in json.loads((ROOT / "assets/euclidean/layouts.json").read_text())["layouts"]}

    def test_missing_font_file_and_checksum(self):
        fonts = self.work / "assets/fonts"
        fonts.mkdir(parents=True)
        entry = json.loads((ROOT / "assets/fonts/fonts.json").read_text())[0]
        (fonts / "fonts.json").write_text(json.dumps([entry]))
        element = self.euclidean.catalogue()["euclidean-cover"]["elements"][1]
        with patch.object(self.euclidean, "ROOT", self.work):
            with self.assertRaisesRegex(renderer.ValidationError, "missing bundled font;"):
                self.euclidean.font_asset(element)
            (fonts / entry["file"]).write_bytes(b"invalid")
            with self.assertRaisesRegex(renderer.ValidationError, "checksum mismatch"):
                self.euclidean.font_asset(element)

    def test_unsupported_glyph_is_not_substituted(self):
        self.content["slides"][0]["body"]["fields"]["caption"] = "\U0001f600"
        self.invalid("lacks glyphs")

    def test_hash_and_approval_are_still_bound_to_theme(self):
        approval = {"status": "approved", "outline_sha256": hashlib.sha256(self.raw).hexdigest(),
                    "outline_version": self.outline["version"], "user_confirmation": "Approved in test only"}
        renderer.validate_approval(approval, self.outline, self.raw)
        with self.assertRaisesRegex(renderer.ValidationError, "renewed approval"):
            renderer.validate_approval(approval, self.outline, self.raw + b"\n")
        outline = self.work / "outline.json"
        outline.write_bytes(self.raw)
        result = subprocess.run([sys.executable, "-S", str(SCRIPT), "outline-hash", "--outline", str(outline)], capture_output=True, text=True)
        self.assertEqual(result.stdout.strip(), approval["outline_sha256"])
        approval_file = self.work / "approval.json"
        approval_file.write_text(json.dumps(approval))
        content_file = self.work / "content.json"
        content_file.write_text(json.dumps(self.content))
        result = self.cli("render", "--outline", outline, "--approval", approval_file,
                          "--content", content_file, "--output-dir", self.work / "approved")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(len(Presentation(self.work / "approved/presentation.pptx").slides), 19)
        outline.write_bytes(self.raw + b"\n")
        result = self.cli("render", "--outline", outline, "--approval", approval_file,
                          "--content", content_file, "--output-dir", self.work / "stale")
        self.assertEqual(result.returncode, 2)
        self.assertFalse((self.work / "stale").exists())

    def test_speech_transition_source_validation_preserved(self):
        self.content["slides"][0]["source_ids"] = ["Missing"]
        self.invalid("unresolved source reference")
        self.content = deepcopy(self.sample_content)
        self.content["slides"][0]["speech"] = ""
        self.invalid("speech: must not be empty")
        self.content = deepcopy(self.sample_content)
        self.content["slides"][0]["transition"] = ""
        self.invalid("transition: must not be empty")

    def test_linear_burn_and_normal_opacity(self):
        element = self.euclidean.catalogue()["euclidean-cover"]["elements"][0]
        paint, opacity = self.euclidean.effective_paint(element, "#FCF7E9")
        self.assertEqual(paint["color"], "#888278")
        self.assertEqual(opacity, 1)

    def test_unknown_and_missing_images(self):
        self.content["slides"][0]["body"]["images"] = {"unknown": "image.png"}
        self.invalid("undeclared image slot")
        self.content = deepcopy(self.sample_content)
        self.content["slides"][4]["body"]["images"] = {"n261": "missing-image.png"}
        self.invalid("local image not found")

    def test_optional_image_keeps_source_geometry_and_z_order(self):
        file = self.work / "image.png"
        Image.new("RGB", (100, 100), "red").save(file)
        self.content["slides"][4]["body"]["images"] = {"n261": str(file)}
        renderer.validate_content(self.content, self.outline)
        deck = Presentation(BytesIO(renderer.build_presentation(self.outline, self.content)))
        shapes = list(deck.slides[4].shapes)
        index = next(i for i, shape in enumerate(shapes) if shape.name == "n261")
        image = shapes[index + 1]
        self.assertEqual(image.name, "n261-image")
        self.assertEqual(image.width, image.height)
        self.assertGreaterEqual(image.left, shapes[index].left)
        self.assertLessEqual(image.left + image.width, shapes[index].left + shapes[index].width)
        self.assertEqual(image.top, shapes[index].top)

    def test_source_preview_is_labeled_and_uses_bundled_fonts(self):
        from xml.etree import ElementTree
        preview = self.euclidean.preview_svg(self.content["slides"][0])
        root = ElementTree.fromstring(preview)
        self.assertEqual(root.attrib["viewBox"], "0 0 1920 1080")
        self.assertIn("not a native PowerPoint rendering", preview)
        self.assertIn("data:font/ttf;base64,", preview)
        self.assertIn("26 JUNE 2024", preview)

    def test_design_credit_is_nonvisual_and_not_a_research_reference(self):
        sources_before = deepcopy(self.content["sources"])
        speech, sources = renderer.markdown_documents(self.outline, self.content, True)
        for document in (speech, sources):
            self.assertIn(self.euclidean.DESIGN_CREDIT, document)
            self.assertIn("not a research source", document)
        deck = Presentation(BytesIO(renderer.build_presentation(self.outline, self.content)))
        for slide in deck.slides:
            self.assertIn(self.euclidean.DESIGN_CREDIT, slide.notes_slide.notes_text_frame.text)
            visible = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            self.assertNotIn("Nicholas Jitkoff", visible)
            self.assertNotIn("creativecommons.org", visible)
        self.assertEqual(self.content["sources"], sources_before)

    def test_ordered_spacing_and_center_alignment(self):
        deck = Presentation(BytesIO(renderer.build_presentation(self.outline, self.content)))
        agenda = next(shape for shape in deck.slides[2].shapes if shape.name == "n247")
        first = agenda.text_frame.paragraphs[0]
        self.assertEqual(first.space_after.pt, 20)
        self.assertAlmostEqual(first.line_spacing.pt, 26.5, places=2)
        statement = next(shape for shape in deck.slides[1].shapes if shape.name == "n241")
        self.assertEqual(statement.text_frame.paragraphs[0]._p.pPr.get("marL"), "0")

    def test_euclidean_theme_and_legacy_theme_isolation(self):
        from xml.etree import ElementTree
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        payload = renderer.build_presentation(self.outline, self.content)
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            theme = ElementTree.fromstring(archive.read("ppt/theme/theme1.xml"))
        self.assertEqual(theme.get("name"), "Euclidean")
        self.assertEqual(theme.find(".//a:majorFont/a:latin", ns).get("typeface"), "Source Serif Pro")
        self.assertEqual(theme.find(".//a:minorFont/a:latin", ns).get("typeface"), "Source Sans Pro")
        colors = {
            "dk1": "16161D", "lt1": "FCF7E9", "dk2": "88888C", "lt2": "FEFBF4",
            "accent1": "D42A20", "accent2": "0E638E", "accent3": "0E8E51",
            "accent4": "FAC22B", "accent5": "E77625", "accent6": "67378D",
        }
        for name, value in colors.items():
            self.assertEqual(theme.find(".//a:clrScheme/a:" + name + "/a:srgbClr", ns).get("val"), value)
        self.assertIsNone(theme.find(".//a:outerShdw", ns))
        old_outline, _ = renderer.load_json(ROOT / "assets/template-outline.json")
        old_content, _ = renderer.load_json(ROOT / "assets/template-deck.json")
        with zipfile.ZipFile(BytesIO(renderer.build_presentation(old_outline, old_content))) as archive:
            legacy_theme = archive.read("ppt/theme/theme1.xml")
        default = BytesIO()
        Presentation().save(default)
        with zipfile.ZipFile(default) as archive:
            self.assertEqual(legacy_theme, archive.read("ppt/theme/theme1.xml"))


if __name__ == "__main__":
    unittest.main()
