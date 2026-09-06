"""Native Euclidean source-layout rendering; no system-font substitution or resizing."""

from functools import lru_cache
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import struct


ROOT = Path(__file__).resolve().parents[1]
EMU_PER_PX = 6350
DESIGN_CREDIT = (
    'Adapted from “Euclidean” by Nicholas Jitkoff, '
    "https://www.figma.com/community/file/1387141600584938870/euclidean ; "
    "CC BY 4.0 https://creativecommons.org/licenses/by/4.0/ ; "
    "adapted to editable PowerPoint."
)


def fail(path, message):
    raise ValueError(f"{path}: {message}")


def catalogue():
    path = ROOT / "assets/euclidean/layouts.json"
    if not path.is_file():
        fail(str(path), "missing Euclidean catalogue; restore the packaged assets/euclidean directory")
    return {item["id"]: item for item in json.loads(path.read_text(encoding="utf-8"))["layouts"]}


def image_slots(layout):
    return [e["id"] for e in layout["elements"] if e["kind"] == "rounded_rectangle"]


def catalog_document():
    document = json.loads((ROOT / "assets/euclidean/layouts.json").read_text(encoding="utf-8"))
    for layout in document["layouts"]:
        layout["image_slots"] = image_slots(layout)
    document["image_policy"] = "Optional local images replace named panels, contain-fit without moving text. Divider background imagery is not supported."
    return document


def font_asset(element):
    path = ROOT / "assets/fonts/fonts.json"
    if not path.is_file():
        fail(str(path), "missing bundled font manifest; restore assets/fonts (no font substitution is allowed)")
    font = element["font"]
    entries = json.loads(path.read_text(encoding="utf-8"))
    for entry in entries:
        if entry["family"] == font["family"] and entry["style"] == font["style"]:
            file = path.parent / entry["file"]
            if not file.is_file():
                fail(str(file), "missing bundled font; restore assets/fonts before rendering")
            if hashlib.sha256(file.read_bytes()).hexdigest() != entry["sha256"]:
                fail(str(file), "bundled font checksum mismatch; restore the original font")
            return file
    fail(str(path), f"missing bundled font {font['family']} {font['style']}; restore the complete font package")


@lru_cache(maxsize=64)
def load_font(file, size, checksum):
    try:
        from PIL import ImageFont
    except ImportError:
        fail("fonts", f"Pillow is missing; install dependencies from {ROOT / 'requirements.txt'}")
    try:
        return ImageFont.truetype(str(file), round(size * 4))
    except OSError as error:
        fail(str(file), f"cannot load bundled font: {error}")


@lru_cache(maxsize=8)
def supported_characters(file, checksum):
    """Read Unicode cmap tables without introducing a fontTools dependency."""
    data = file.read_bytes()
    count = struct.unpack_from(">H", data, 4)[0]
    tables = {}
    for index in range(count):
        tag, _, offset, _ = struct.unpack_from(">4sIII", data, 12 + index * 16)
        tables[tag] = offset
    cmap = tables[b"cmap"]
    codes = set()
    for index in range(struct.unpack_from(">H", data, cmap + 2)[0]):
        platform, encoding, offset = struct.unpack_from(">HHI", data, cmap + 4 + index * 8)
        if platform != 0 and not (platform == 3 and encoding in (1, 10)):
            continue
        start = cmap + offset
        format_id = struct.unpack_from(">H", data, start)[0]
        if format_id == 12:
            for group in range(struct.unpack_from(">I", data, start + 12)[0]):
                first, last, glyph = struct.unpack_from(">III", data, start + 16 + group * 12)
                codes.update(range(first + (glyph == 0), last + 1))
        elif format_id == 4:
            segments = struct.unpack_from(">H", data, start + 6)[0] // 2
            ends = start + 14
            starts = ends + segments * 2 + 2
            deltas = starts + segments * 2
            offsets = deltas + segments * 2
            for segment in range(segments):
                first = struct.unpack_from(">H", data, starts + segment * 2)[0]
                last = struct.unpack_from(">H", data, ends + segment * 2)[0]
                delta = struct.unpack_from(">h", data, deltas + segment * 2)[0]
                shift = struct.unpack_from(">H", data, offsets + segment * 2)[0]
                for code in range(first, min(last, 65534) + 1):
                    glyph = (code + delta) & 65535
                    if shift:
                        glyph = struct.unpack_from(">H", data, offsets + segment * 2 + shift + (code - first) * 2)[0]
                        glyph = (glyph + delta) & 65535 if glyph else 0
                    if glyph:
                        codes.add(code)
    return codes


def spacing(element):
    value = element["letter_spacing"]
    return value["value"] * element["font"]["size"] / 100 if value["units"] == "PERCENT" else value["value"]


def line_height(element):
    value, size = element["line_height"], element["font"]["size"]
    if value["units"] == "RAW":
        return value["value"] * size
    if value["units"] == "PERCENT":
        return value["value"] * size / 100
    return value["value"]


def width_of(value, font, tracking):
    return font.getlength(value) / 4 + max(0, len(value) - 1) * tracking


def display_text(value, element):
    return value.upper() if element["text_case"] == "UPPER" else value


def text_layout(element, value, path):
    """Measure actual fonts, preserving source baseline records only for source text."""
    text(value, path, 10000)
    file = font_asset(element)
    checksum = hashlib.sha256(file.read_bytes()).hexdigest()
    font = load_font(file, element["font"]["size"], checksum)
    tracking = spacing(element)
    value = display_text(value, element)
    unsupported = sorted({c for c in value if c != "\n" and ord(c) not in supported_characters(file, checksum)})
    if unsupported:
        fail(path, "bundled font lacks glyphs " + ", ".join(f"U+{ord(c):04X}" for c in unsupported) + "; choose supported text (font substitution is disabled)")
    original = display_text(element["text"], element)
    ordered = element["paragraphs"][0]["type"] == "ORDERED_LIST"
    indent = element["font"]["size"] * 1.5 if ordered else 0
    capacity = element["size"][0] - indent
    height = line_height(element)
    ascent, descent = (n / 4 for n in font.getmetrics())
    source_baseline = element["baselines"][0]
    # Figma's natural-height snapshot differs from FreeType hhea metrics. Its
    # 100% line boxes use that snapshot; explicit larger line heights do not.
    natural = source_baseline["height"]
    advance = natural if height <= element["font"]["size"] else height
    rows = []
    if value == original and element.get("baselines"):
        for index, baseline in enumerate(element["baselines"]):
            line = value[baseline["start"]:baseline["end"]].strip("\n ")
            rows.append({"text": line, "x": baseline["x"], "baseline": baseline["y"],
                         "paragraph": index if ordered else 0, "first": ordered or index == 0,
                         "marker_x": baseline.get("marker_x")})
        # Figma source ranges include paragraph separators; splitting list entries
        # explicitly prevents separators from becoming visible text.
        if ordered:
            paragraphs = value.split("\n")
            if len(paragraphs) == len(rows):
                for row, paragraph in zip(rows, paragraphs):
                    row["text"] = paragraph
        for row in rows:
            if width_of(row["text"], font, tracking) + row["x"] > element["size"][0] + 1:
                fail(path, "source baseline exceeds its box with the bundled font; verify the original font assets")
        return {"rows": rows, "font": font, "file": file, "ascent": ascent,
                "advance": advance, "ordered": ordered, "source": True}
    for paragraph_index, paragraph in enumerate(value.split("\n")):
        words = re.findall(r"[^ \t\r\n]+", paragraph)
        current, lines = "", []
        for word in words:
            if width_of(word, font, tracking) > capacity + 0.5:
                fail(path, f"word exceeds {capacity:g}px at fixed {element['font']['size'] / 2:g}pt; shorten text or choose another layout")
            candidate = current + (" " if current else "") + word
            if current and width_of(candidate, font, tracking) > capacity + 0.5:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
        for index, line in enumerate(lines):
            baseline = (rows[-1]["baseline"] + advance if rows else source_baseline["y"])
            if index == 0 and rows:
                baseline += element["list_spacing"]
            rows.append({"text": line, "x": indent, "baseline": baseline,
                         "paragraph": paragraph_index, "first": index == 0,
                         "marker_x": indent - width_of(f"{paragraph_index + 1}.", font, tracking)
                         - element["font"]["size"] / 3 if ordered else None})
    required = len(rows) * advance + max(0, len(value.split("\n")) - 1) * element["list_spacing"]
    if required > element["size"][1] + 0.5:
        fail(path, f"content needs {required:.1f}px height; slot is {element['size'][1]:g}px at fixed {element['font']['size'] / 2:g}pt; shorten text or choose another layout")
    extra = max(0, element["size"][1] - required)
    offset = extra / 2 if element["vertical_align"] == "CENTER" else extra if element["vertical_align"] == "BOTTOM" else 0
    for row in rows:
        row["baseline"] += offset
        bounds = font.getbbox(row["text"], anchor="ls")
        if row["baseline"] + bounds[3] / 4 > element["size"][1] + .5:
            fail(path, "glyph descenders exceed the fixed source text box; shorten text or choose another layout")
    return {"rows": rows, "font": font, "file": file, "ascent": ascent,
            "advance": advance, "ordered": ordered, "source": False}


def validate_title(slide, path):
    layout = catalogue()[slide["layout"]]
    element = next(e for e in layout["elements"] if e.get("slot") == "title")
    text_layout(element, slide["title"], path + ".title")


def validate_body(slide, path):
    layout = catalogue()[slide["layout"]]
    body = slide["body"]
    fields(body, {"fields"} | ({"images"} if isinstance(body, dict) and "images" in body else set()), path)
    fields(body["fields"], set(layout["fields"]), path + ".fields")
    for element in layout["elements"]:
        slot = element.get("slot")
        if element["kind"] == "text":
            value = slide["title"] if slot == "title" else body["fields"][slot] if slot else element["text"]
            text_layout(element, value, path + ".fields." + str(slot))
    images = body.get("images", {})
    if not isinstance(images, dict):
        fail(path + ".images", "expected an object keyed by catalogue image slot IDs")
    for slot, value in images.items():
        if slot not in image_slots(layout):
            fail(path + ".images", f"undeclared image slot {slot}; use only catalogue image_slots")
        text(value, path + ".images." + slot, 4096)
        file = Path(value)
        if not file.is_file():
            fail(path + ".images." + slot, f"local image not found: {file}")
        from PIL import Image
        try:
            with Image.open(file) as image:
                if image.format not in ("PNG", "JPEG"):
                    fail(path + ".images." + slot, "use a local PNG or JPEG image")
                image.verify()
        except (OSError, ValueError) as error:
            fail(path + ".images." + slot, f"invalid local image: {error}")


def px(value):
    return round(value * EMU_PER_PX)


def point(matrix, x, y):
    a, b, c, d, tx, ty = matrix
    return a * x + c * y + tx, b * x + d * y + ty


def xml(tag, **attributes):
    from pptx.oxml.xmlchemy import OxmlElement
    result = OxmlElement("a:" + tag)
    for key, value in attributes.items():
        result.set(key, str(value))
    return result


def color_fill(fill, paint, opacity=1):
    from pptx.dml.color import RGBColor
    if paint is None:
        fill.background()
        return
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(paint["color"].lstrip("#"))
    alpha = opacity * paint.get("opacity", 1)
    if alpha < 1:
        fill._xPr.solidFill.srgbClr.append(xml("alpha", val=round(alpha * 100000)))


def effective_paint(element, background):
    paint = element.get("fill")
    opacity = element["opacity"]
    if paint and element["blend"] == "LINEAR_BURN":
        source = [int(paint["color"][i:i + 2], 16) for i in (1, 3, 5)]
        back = [int(background[i:i + 2], 16) for i in (1, 3, 5)]
        alpha = opacity * paint.get("opacity", 1)
        channels = [round(b * (1 - alpha) + max(0, b + s - 255) * alpha) for s, b in zip(source, back)]
        return {"color": "#" + "".join(f"{c:02X}" for c in channels), "opacity": 1}, 1
    return paint, opacity


def custom_geometry(shape, element):
    geometry = xml("custGeom")
    for name in ("avLst", "gdLst", "ahLst", "cxnLst"):
        geometry.append(xml(name))
    geometry.append(xml("rect", l="0", t="0", r="r", b="b"))
    paths = xml("pathLst")
    path = xml("path", w=px(element["size"][0]), h=px(element["size"][1]), stroke="0", extrusionOk="0")
    commands = {"M": "moveTo", "L": "lnTo", "Q": "quadBezTo", "C": "cubicBezTo", "Z": "close"}
    for command in element["path"]:
        node = xml(commands[command[0]])
        for i in range(1, len(command), 2):
            node.append(xml("pt", x=px(command[i]), y=px(command[i + 1])))
        path.append(node)
    paths.append(path)
    geometry.append(paths)
    properties = shape._element.spPr
    properties.remove(properties.prstGeom)
    properties.insert(1, geometry)


def add_text(slide, element, value, background, language):
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt
    x, y = point(element["transform"], 0, 0)
    w, h = element["size"]
    measured = text_layout(element, value, element["id"])
    rows = measured["rows"]
    # A zero-height quote is positioned by its source baseline, not discarded.
    shape = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = frame.margin_bottom = 0
    frame.margin_top = px(rows[0]["baseline"] - measured["ascent"])
    frame._txBody.bodyPr.set("vertOverflow", "overflow")
    frame._txBody.bodyPr.set("horzOverflow", "overflow")
    paint, opacity = effective_paint(element, background)
    style = element["font"]["style"]
    family = element["font"]["family"]
    face = family if style in ("Regular", "Bold") else family + " " + ("Semibold" if style == "SemiBold" else style)
    for index, row in enumerate(rows):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = {"LEFT": PP_ALIGN.LEFT, "CENTER": PP_ALIGN.CENTER, "RIGHT": PP_ALIGN.RIGHT}[element["align"]]
        paragraph.space_before = Pt(0)
        following = rows[index + 1] if index + 1 < len(rows) else None
        gap = element["list_spacing"] if following and following["first"] else 0
        paragraph.space_after = Pt(gap / 2)
        paragraph.line_spacing = Pt((following["baseline"] - row["baseline"] - gap) / 2 if following else measured["advance"] / 2)
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(px(row["x"] if element["align"] == "LEFT" else 0)))
        if measured["ordered"] and row["first"]:
            properties.set("indent", str(px(row["marker_x"] - row["x"])))
            properties.append(xml("buFont", typeface=face))
            properties.append(xml("buAutoNum", type="arabicPeriod", startAt=row["paragraph"] + 1))
        else:
            properties.append(xml("buNone"))
        run = paragraph.add_run()
        run.text = row["text"]
        run.font.name = face
        run.font.size = Pt(element["font"]["size"] / 2)
        run.font.bold = style == "Bold"
        run.font.italic = False
        rpr = run._r.get_or_add_rPr()
        rpr.set("spc", str(round(spacing(element) * 50)))
        rpr.set("kern", "0")
        rpr.set("lang", language)
        for name in ("ea", "cs"):
            rpr.append(xml(name, typeface=face))
        color_fill(run.font.fill, paint, opacity)
    return shape


def configure_theme(deck):
    """Make subsequent native edits inherit Euclidean rather than Office styling."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    theme_part = deck.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    root.set("name", "Euclidean")
    colors = {
        "dk1": "16161D", "lt1": "FCF7E9", "dk2": "88888C", "lt2": "FEFBF4",
        "accent1": "D42A20", "accent2": "0E638E", "accent3": "0E8E51",
        "accent4": "FAC22B", "accent5": "E77625", "accent6": "67378D",
        "hlink": "0E638E", "folHlink": "67378D",
    }
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", "Euclidean")
    for name, color in colors.items():
        node = scheme.find("a:" + name, ns)
        for child in list(node):
            node.remove(child)
        node.append(xml("srgbClr", val=color))
    fonts = root.find(".//a:fontScheme", ns)
    fonts.set("name", "Euclidean")
    for name, family in (("majorFont", "Source Serif Pro"), ("minorFont", "Source Sans Pro")):
        node = fonts.find("a:" + name, ns)
        for child in list(node):
            node.remove(child)
        for tag in ("latin", "ea", "cs"):
            node.append(xml(tag, typeface=family))
    root.find(".//a:fmtScheme", ns).set("name", "Euclidean")
    for style in root.findall(".//a:effectStyle", ns):
        for child in list(style):
            style.remove(child)
        style.append(xml("effectLst"))
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def build(outline, content, placeholder, notes_text, Presentation):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.util import Pt
    deck = Presentation()
    configure_theme(deck)
    deck.slide_width, deck.slide_height = px(1920), px(1080)
    deck.core_properties.title = outline["title"]
    deck.core_properties.subject = f"Euclidean; Mode: {outline['mode']}"
    deck.core_properties.language = outline["language"]
    deck.core_properties.author = "Presentation generator"
    deck.core_properties.keywords = "DEMO, PLACEHOLDER, NOT APPROVED" if placeholder else "euclidean, approved outline"
    layouts = catalogue()
    sources = {source["id"]: source for source in content["sources"]}
    for data, planned in zip(content["slides"], outline["slides"]):
        layout = layouts[data["layout"]]
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(layout["background"].lstrip("#"))
        for element in layout["elements"]:
            kind = element["kind"]
            x, y = point(element["transform"], 0, 0)
            w, h = element["size"]
            if kind == "text":
                slot = element["slot"]
                value = data["title"] if slot == "title" else data["body"]["fields"][slot] if slot else element["text"]
                shape = add_text(slide, element, value, layout["background"], outline["language"])
            elif kind == "line":
                end = point(element["transform"], w, h)
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x), px(y), px(end[0]), px(end[1]))
            else:
                shape = slide.shapes.add_shape(MSO_SHAPE.OVAL if kind == "ellipse" else MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
                if kind == "vector":
                    custom_geometry(shape, element)
                color_fill(shape.fill, element["fill"], element["opacity"])
            if kind != "text":
                color_fill(shape.line.fill, element["stroke"], element["opacity"])
                if element["stroke"]:
                    shape.line.width = Pt(element["stroke_width"] / 2)
            # python-pptx's default shape style references an Office shadow.
            # Source selection/UI shadows are not part of the Euclidean artwork.
            for style in shape._element.xpath("./p:style"):
                shape._element.remove(style)
            shape._element.spPr.append(xml("effectLst"))
            shape.name = element["id"]
            if kind == "text":
                shape._element._nvXxPr.cNvPr.set("descr", display_text(value, element))
            if element["id"] in data["body"].get("images", {}):
                from PIL import Image
                image_path = data["body"]["images"][element["id"]]
                with Image.open(image_path) as image:
                    iw, ih = image.size
                scale = min(w / iw, h / ih)
                picture = slide.shapes.add_picture(image_path, px(x + (w - iw * scale) / 2), px(y + (h - ih * scale) / 2), width=px(iw * scale), height=px(ih * scale))
                picture.name = element["id"] + "-image"
        slide.notes_slide.notes_text_frame.text = ("DEMO — source-layout example, not approved.\n\n" if placeholder else "") + notes_text(data, planned, outline, sources, placeholder)
    result = BytesIO()
    deck.save(result)
    return result.getvalue()


def preview_svg(data):
    """Source-layout preview, NOT a screenshot or native PowerPoint verification."""
    import base64
    from html import escape
    layout = catalogue()[data["layout"]]
    content, fonts = [], {}
    for element in layout["elements"]:
        kind = element["kind"]
        matrix = "matrix(" + " ".join(str(n) for n in element["transform"]) + ")"
        paint, opacity = effective_paint(element, layout["background"])
        fill = paint["color"] if paint else "none"
        alpha = opacity * (paint.get("opacity", 1) if paint else 1)
        stroke = element["stroke"]
        attrs = f'transform="{matrix}" fill="{fill}" fill-opacity="{alpha}"'
        if stroke:
            attrs += f' stroke="{stroke["color"]}" stroke-opacity="{element["opacity"] * stroke.get("opacity", 1)}" stroke-width="{element["stroke_width"]}"'
        w, h = element["size"]
        if kind == "text":
            slot = element["slot"]
            value = data["title"] if slot == "title" else data["body"]["fields"][slot] if slot else element["text"]
            measured = text_layout(element, value, element["id"])
            font_key = str(measured["file"])
            if font_key not in fonts:
                fonts[font_key] = "font" + str(len(fonts))
            attrs += f' font-family="{fonts[font_key]}" font-size="{element["font"]["size"]}" letter-spacing="{spacing(element)}"'
            lines = []
            for row in measured["rows"]:
                x = row["x"]
                if element["align"] == "CENTER":
                    x = (w - width_of(row["text"], measured["font"], spacing(element))) / 2
                elif element["align"] == "RIGHT":
                    x = w - width_of(row["text"], measured["font"], spacing(element))
                lines.append(f'<text x="{x}" y="{row["baseline"]}" {attrs}>{escape(row["text"])}</text>')
                if measured["ordered"] and row["first"]:
                    lines.append(f'<text x="{row["marker_x"]}" y="{row["baseline"]}" {attrs}>{row["paragraph"] + 1}.</text>')
            content.extend(lines)
        elif kind == "line":
            content.append(f'<line x1="0" y1="0" x2="{w}" y2="{h}" {attrs}/>')
        elif kind == "vector":
            path = " ".join(" ".join(map(str, command)) for command in element["path"])
            content.append(f'<path d="{path}" {attrs}/>')
        elif kind == "ellipse":
            content.append(f'<ellipse cx="{w / 2}" cy="{h / 2}" rx="{w / 2}" ry="{h / 2}" {attrs}/>')
        else:
            content.append(f'<rect width="{w}" height="{h}" {attrs}/>')
        if element["id"] in data["body"].get("images", {}):
            file = Path(data["body"]["images"][element["id"]])
            mime = "image/png" if file.suffix.lower() == ".png" else "image/jpeg"
            uri = f"data:{mime};base64," + base64.b64encode(file.read_bytes()).decode()
            content.append(f'<image width="{w}" height="{h}" transform="{matrix}" href="{uri}" preserveAspectRatio="xMidYMid meet"/>')
    css = "".join(
        "@font-face{font-family:" + name + ";src:url(data:font/ttf;base64," +
        base64.b64encode(Path(file).read_bytes()).decode() + ") format('truetype');}"
        for file, name in fonts.items()
    )
    return (
        '<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" viewBox="0 0 1920 1080">'
        '<title>Source-layout preview — not a native PowerPoint rendering</title>'
        f'<style>{css}</style><rect width="1920" height="1080" fill="{layout["background"]}"/>'
        + "".join(content) + "</svg>"
    )
