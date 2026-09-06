#!/usr/bin/env python3
"""Validate approved presentation inputs and create editable, local PowerPoint decks."""

from __future__ import annotations

import argparse
from datetime import date
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import re
import sys
import unicodedata


ROOT = Path(__file__).resolve().parents[1]
LAYOUTS = {"cover", "bullets", "comparison", "metrics", "table", "chart", "closing"}
WHITE, DARK, BLUE, PANEL, MUTED = "FFFFFF", "172B4D", "2563EB", "EFF6FF", "52657A"
WIDTH, HEIGHT = 13.333, 7.5
FONT = "Arial"
PLACEHOLDER = "PLACEHOLDER / ПРИМЕР — не утверждено; не фактические данные."


class ValidationError(ValueError):
    """A user-facing, actionable input validation failure."""


def fail(path: str, message: str) -> None:
    raise ValidationError(f"{path}: {message}")


def fields(value, required: set[str], path: str) -> dict:
    if not isinstance(value, dict):
        fail(path, "expected an object")
    missing = required - value.keys()
    extra = value.keys() - required
    if missing:
        fail(path, f"missing fields: {', '.join(sorted(missing))}")
    if extra:
        fail(path, f"unknown fields: {', '.join(sorted(extra))}")
    return value


def text(value, path: str, maximum: int = 2000, empty: bool = False) -> str:
    if not isinstance(value, str):
        fail(path, "expected a string")
    if not empty and not value.strip():
        fail(path, "must not be empty")
    if len(value) > maximum:
        fail(path, f"must contain at most {maximum} characters")
    if any(unicodedata.category(c) == "Cc" and c != "\n" for c in value):
        fail(path, "control characters (including tabs) are not supported")
    if any(unicodedata.category(c) == "Cs" or c in "\ufffe\uffff" for c in value):
        fail(path, "contains a character unsupported by XML")
    return value


def identifier(value, path: str) -> str:
    text(value, path, 32)
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{0,31}", value):
        fail(path, "use 1–32 ASCII letters, digits, '_' or '-', starting with a letter")
    return value


def array(value, path: str, minimum: int, maximum: int) -> list:
    if not isinstance(value, list) or not minimum <= len(value) <= maximum:
        fail(path, f"expected an array of {minimum}–{maximum} items")
    return value


def number(value, path: str, minimum: float, maximum: float) -> float:
    if type(value) not in (int, float):
        fail(path, "expected a finite number, not a boolean")
    if not minimum <= value <= maximum or not math.isfinite(value):
        fail(path, f"expected a finite number in [{minimum}, {maximum}]")
    return value


def glyph_width(char: str) -> float:
    """Conservative Arial width estimate in em; no platform font dependencies."""
    if unicodedata.combining(char):
        return 0.0
    if char.isspace():
        return 0.34
    if char in "MW@%ЖШЩЮФЫжшщюфы":
        return 1.02
    if char in "ilIjtfr.,:;'!|()[]":
        return 0.38
    if unicodedata.east_asian_width(char) in ("F", "W"):
        return 1.1
    if char.isupper():
        return 0.79
    if "\u0400" <= char <= "\u04ff":
        return 0.72
    return 0.65


def wrapped(value: str, width: float, size: int, max_lines: int, path: str) -> str:
    text(value, path)
    capacity = width * 72 / (size * 1.08)
    lines = []
    for paragraph in value.split("\n"):
        line, used = [], 0.0
        for word in paragraph.split():
            length = sum(glyph_width(c) for c in word)
            if length > capacity:
                fail(path, f"a word exceeds the {width:g}in text area at {size}pt; shorten it")
            space = glyph_width(" ") if line else 0
            if line and used + space + length > capacity:
                lines.append(" ".join(line))
                line, used, space = [], 0.0, 0.0
            line.append(word)
            used += space + length
        lines.append(" ".join(line))
    if len(lines) > max_lines:
        fail(path, f"content needs {len(lines)} lines; limit is {max_lines} at {size}pt; split or shorten it")
    return "\n".join(lines)


def fit(value, path: str, width: float, size: int, lines: int, empty=False) -> str:
    text(value, path, empty=empty)
    if not value and empty:
        return value
    return wrapped(value, width, size, lines, path)


def unique(values: list, path: str) -> None:
    if len(set(values)) != len(values):
        fail(path, "duplicate IDs or names are not permitted")


def reject_json_constant(value):
    fail("JSON", f"{value} is not a finite JSON number")


def unique_json_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            fail("JSON", f"duplicate object key: {key}")
        result[key] = value
    return result


def load_json(path: Path):
    raw = path.read_bytes()
    if len(raw) > 5_000_000:
        fail(str(path), "JSON file exceeds 5 MB")
    return json.loads(
        raw.decode("utf-8"), parse_constant=reject_json_constant,
        object_pairs_hook=unique_json_object,
    ), raw


def validate_outline(outline: dict) -> None:
    fields(outline, {"version", "title", "mode", "language", "audience",
                     "duration_minutes", "slides"} | ({"theme"} if isinstance(outline, dict) and "theme" in outline else set()), "outline")
    theme = outline.get("theme", "light-business")
    if theme not in ("euclidean", "light-business"):
        fail("outline.theme", "expected euclidean or light-business")
    layouts = euclidean().catalogue() if theme == "euclidean" else LAYOUTS
    if type(outline["version"]) is not int or outline["version"] < 1:
        fail("outline.version", "must be a positive integer")
    text(outline["title"], "outline.title", 255)
    text(outline["language"], "outline.language", 35)
    text(outline["audience"], "outline.audience", 500)
    if outline["mode"] not in ("research", "proposal", "report"):
        fail("outline.mode", "must be research, proposal, or report")
    number(outline["duration_minutes"], "outline.duration_minutes", 0, 360)
    if outline["duration_minutes"] <= 0:
        fail("outline.duration_minutes", "must be positive")
    slides = array(outline["slides"], "outline.slides", 1, 100)
    ids, total = [], 0.0
    for i, slide in enumerate(slides):
        path = f"outline.slides[{i}]"
        fields(slide, {"id", "title", "summary", "layout", "seconds"}, path)
        ids.append(identifier(slide["id"], path + ".id"))
        text(slide["summary"], path + ".summary", 2000)
        if not isinstance(slide["layout"], str) or slide["layout"] not in layouts:
            fail(path + ".layout", f"expected one of {', '.join(sorted(layouts))}")
        if theme == "euclidean":
            euclidean().validate_title(slide, path)
        else:
            fit(slide["title"], path + ".title", 11.9, 30, 2)
        total += number(slide["seconds"], path + ".seconds", 0, 21600)
        if slide["seconds"] <= 0:
            fail(path + ".seconds", "must be positive")
    unique(ids, "outline.slides")
    if total > outline["duration_minutes"] * 60 + 1e-8:
        fail("outline.slides.seconds", "total exceeds duration_minutes × 60")


def validate_approval(approval: dict, outline: dict, outline_bytes: bytes) -> None:
    if not isinstance(approval, dict):
        fail("approval", "expected an object")
    for key in ("status", "outline_sha256", "outline_version", "user_confirmation"):
        if key not in approval:
            fail("approval", f"missing field: {key}")
    if approval["status"] != "approved":
        fail("approval.status", "must be 'approved'; explicit user approval is required")
    expected = hashlib.sha256(outline_bytes).hexdigest()
    if approval["outline_sha256"] != expected:
        fail("approval.outline_sha256", "does not match the raw outline bytes; obtain renewed approval")
    if type(approval["outline_version"]) is not int or approval["outline_version"] != outline["version"]:
        fail("approval.outline_version", "does not match outline.version")
    text(approval["user_confirmation"], "approval.user_confirmation", 10000)


def date_string(value, path: str, allow_nd=False) -> None:
    text(value, path, 10)
    if allow_nd and value == "n.d.":
        return
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        fail(path, "expected YYYY-MM-DD" + (" or n.d." if allow_nd else ""))
    try:
        date.fromisoformat(value)
    except ValueError as error:
        raise ValidationError(f"{path}: invalid calendar date") from error


def validate_body(slide: dict, path: str) -> None:
    body, layout = slide["body"], slide["layout"]
    if layout == "cover":
        fields(body, {"subtitle", "presenter"}, path)
        fit(body["subtitle"], path + ".subtitle", 10.9, 26, 3)
        fit(body["presenter"], path + ".presenter", 10.9, 18, 2)
    elif layout == "bullets":
        fields(body, {"items"}, path)
        for i, value in enumerate(array(body["items"], path + ".items", 1, 5)):
            fit(value, f"{path}.items[{i}]", 10.7, 22, 2)
    elif layout == "comparison":
        fields(body, {"left_title", "left_items", "right_title", "right_items"}, path)
        for side in ("left", "right"):
            fit(body[f"{side}_title"], f"{path}.{side}_title", 5.02, 22, 2)
            for i, value in enumerate(array(body[f"{side}_items"], f"{path}.{side}_items", 1, 4)):
                fit(value, f"{path}.{side}_items[{i}]", 4.62, 18, 2)
    elif layout == "metrics":
        fields(body, {"items"}, path)
        items = array(body["items"], path + ".items", 1, 3)
        width = (11.9 - 0.3 * (len(items) - 1)) / len(items) - 0.6
        for i, item in enumerate(items):
            p = f"{path}.items[{i}]"
            fields(item, {"value", "label", "context"}, p)
            fit(item["value"], p + ".value", width, 36, 1)
            fit(item["label"], p + ".label", width, 22, 2)
            fit(item["context"], p + ".context", width, 18, 3)
    elif layout == "table":
        fields(body, {"columns", "rows"}, path)
        columns = array(body["columns"], path + ".columns", 1, 5)
        width = 11.9 / len(columns) - 0.24
        for i, value in enumerate(columns):
            fit(value, f"{path}.columns[{i}]", width, 14, 2)
        rows = array(body["rows"], path + ".rows", 1, 6)
        for i, row in enumerate(rows):
            array(row, f"{path}.rows[{i}]", len(columns), len(columns))
            for j, value in enumerate(row):
                fit(value, f"{path}.rows[{i}][{j}]", width, 14, 2, empty=True)
    elif layout == "chart":
        fields(body, {"categories", "series", "unit"}, path)
        categories = array(body["categories"], path + ".categories", 1, 6)
        for i, value in enumerate(categories):
            fit(value, f"{path}.categories[{i}]", 1.35, 18, 1)
        unique(categories, path + ".categories")
        text(body["unit"], path + ".unit")
        fit(f"Unit / Единица: {body['unit']}", path + ".unit", 11.9, 18, 1)
        series = array(body["series"], path + ".series", 1, 3)
        for i, item in enumerate(series):
            p = f"{path}.series[{i}]"
            fields(item, {"name", "values"}, p)
            fit(item["name"], p + ".name", 2.8, 18, 1)
            for j, value in enumerate(array(item["values"], p + ".values", len(categories), len(categories))):
                number(value, f"{p}.values[{j}]", -1e9, 1e9)
                if value != 0 and abs(value) < 1e-9:
                    fail(f"{p}.values[{j}]", "nonzero chart magnitudes must be at least 1e-9")
        unique([item["name"] for item in series], path + ".series")
    elif layout == "closing":
        fields(body, {"takeaway", "actions"}, path)
        fit(body["takeaway"], path + ".takeaway", 10.7, 26, 3)
        for i, value in enumerate(array(body["actions"], path + ".actions", 1, 3)):
            fit(value, f"{path}.actions[{i}]", 10.5, 20, 2)


def validate_content(content: dict, outline: dict) -> None:
    fields(content, {"slides", "sources"}, "content")
    sources = array(content["sources"], "content.sources", 0, 500)
    ids = []
    for i, source in enumerate(sources):
        path = f"content.sources[{i}]"
        fields(source, {"id", "title", "publisher", "date", "location", "locator",
                        "access", "accessed_at"}, path)
        ids.append(identifier(source["id"], path + ".id"))
        for key in ("title", "publisher", "location", "locator"):
            text(source[key], path + "." + key, 2000)
        if source["access"] not in ("public", "local", "notebooklm", "restricted"):
            fail(path + ".access", "must be public, local, notebooklm, or restricted")
        date_string(source["date"], path + ".date", allow_nd=True)
        date_string(source["accessed_at"], path + ".accessed_at")
    unique(ids, "content.sources")
    slides = array(content["slides"], "content.slides", len(outline["slides"]), len(outline["slides"]))
    for i, (slide, planned) in enumerate(zip(slides, outline["slides"])):
        path = f"content.slides[{i}]"
        fields(slide, {"id", "title", "layout", "body", "speech", "transition", "source_ids"}, path)
        for key in ("id", "title", "layout"):
            if slide[key] != planned[key]:
                fail(path + "." + key, f"must exactly match ordered outline slide {planned['id']}")
        text(slide["speech"], path + ".speech", 30000)
        text(slide["transition"], path + ".transition", 2000, empty=i == len(slides) - 1)
        refs = array(slide["source_ids"], path + ".source_ids", 0, 12)
        for j, ref in enumerate(refs):
            identifier(ref, f"{path}.source_ids[{j}]")
            if ref not in ids:
                fail(path + ".source_ids", f"unresolved source reference: {ref}")
        unique(refs, path + ".source_ids")
        if outline.get("theme") == "euclidean":
            euclidean().validate_body(slide, path + ".body")
        else:
            fit(source_footer(refs), path + ".source_ids (footer)", 10.7, 12, 1)
            validate_body(slide, path + ".body")


def source_footer(refs: list[str]) -> str:
    return "Sources / Источники: " + (", ".join(refs) if refs else "—")


def source_detail(source: dict) -> str:
    return (
        f"[{source['id']}] {source['title']}\n"
        f"Publisher / Издатель: {source['publisher']}; date: {source['date']}\n"
        f"Location / Адрес: {source['location']}\n"
        f"Locator / Фрагмент: {source['locator']}\n"
        f"Access: {source['access']}; accessed: {source['accessed_at']}"
    )


def notes_text(slide: dict, planned: dict, outline: dict, sources: dict, placeholder: bool) -> str:
    details = "\n\n".join(source_detail(sources[ref]) for ref in slide["source_ids"])
    return (
        (PLACEHOLDER + "\n\n" if placeholder else "")
        + ("Design credit (not a research source):\n" + euclidean().DESIGN_CREDIT + "\n\n"
           if outline.get("theme") == "euclidean" else "")
        + f"{slide['id']} — {slide['title']}\n"
        f"Language / Язык: {outline['language']}\nAudience / Аудитория: {outline['audience']}\n"
        f"Timing / Время: {planned['seconds']:g} seconds\n\n"
        f"Speech / Речь:\n{slide['speech']}\n\n"
        f"Transition / Переход:\n{slide['transition'] or '—'}\n\n"
        f"Sources / Источники:\n{details or '— (no references supplied)'}"
    )


def load_pptx():
    try:
        from pptx import Presentation
    except ModuleNotFoundError as error:
        raise ValidationError(
            "PowerPoint dependency missing; install with your virtual environment's "
            f"python -m pip install -r \"{ROOT / 'requirements.txt'}\""
        ) from error
    return Presentation


def euclidean():
    # Lazy loading keeps outline-hash independent of rendering/font dependencies,
    # including callers loading this script with importlib rather than as a CLI.
    import importlib.util
    name = "_slides_euclidean"
    if name not in sys.modules:
        spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / "euclidean.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        module.fail, module.text, module.fields = fail, text, fields
        sys.modules[name] = module
    return sys.modules[name]


def literal_chart_workbook(body: dict) -> bytes:
    from xlsxwriter import Workbook

    result = BytesIO()
    # Chart labels are data, even when they start with "=" or resemble a URL.
    with Workbook(result, {"in_memory": True, "strings_to_formulas": False,
                           "strings_to_urls": False}) as workbook:
        sheet = workbook.add_worksheet("Sheet1")
        for row, category in enumerate(body["categories"], start=1):
            sheet.write_string(row, 0, category)
        for column, series in enumerate(body["series"], start=1):
            sheet.write_string(0, column, series["name"])
            for row, value in enumerate(series["values"], start=1):
                sheet.write_number(row, column, value)
    return result.getvalue()


def build_presentation(outline: dict, content: dict, placeholder=False) -> bytes:
    if outline.get("theme") == "euclidean":
        return euclidean().build(outline, content, placeholder, notes_text, load_pptx())
    Presentation = load_pptx()
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(WIDTH), Inches(HEIGHT)
    deck.core_properties.title = outline["title"]
    deck.core_properties.subject = f"Mode: {outline['mode']}"
    deck.core_properties.language = outline["language"]
    deck.core_properties.author = "Presentation generator"
    deck.core_properties.keywords = "PLACEHOLDER, NOT APPROVED" if placeholder else "approved outline"
    def font_style(font, size, color=DARK, bold=False):
        font.name, font.size, font.bold = FONT, Pt(size), bold
        font.color.rgb = RGBColor.from_string(color)

    def fill(shape, color):
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()

    def panel(slide, x, y, w, h, color=PANEL):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        fill(shape, color)
        return shape

    def write_frame(frame, value, size, color=DARK, bold=False):
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        for index, line in enumerate(value.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.15
            font_style(paragraph.font, size, color, bold)
            for run in paragraph.runs:
                font_style(run.font, size, color, bold)
                run._r.get_or_add_rPr().set("lang", outline["language"])

    def box(slide, value, x, y, w, h, size=22, color=DARK, bold=False, lines=2):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.margin_left = frame.margin_right = Inches(0)
        frame.margin_top = frame.margin_bottom = Inches(0)
        formatted = wrapped(value, w, size, lines, "rendered text") if value else ""
        write_frame(frame, formatted, size, color, bold)
        return shape

    sources = {source["id"]: source for source in content["sources"]}
    for i, (data, planned) in enumerate(zip(content["slides"], outline["slides"])):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
        panel(slide, 0.7, 0.4, 0.65, 0.06, BLUE)
        box(slide, data["title"], 0.7, 0.7, 11.9, 1.05, size=30, bold=True)
        layout, body = data["layout"], data["body"]
        if layout == "cover":
            panel(slide, 0.7, 2.05, 11.9, 4.45)
            box(slide, body["subtitle"], 1.2, 2.65, 10.9, 1.65, size=26, lines=3)
            box(slide, body["presenter"], 1.2, 5.05, 10.9, 0.8, size=18, color=MUTED)
        elif layout == "bullets":
            for j, value in enumerate(body["items"]):
                y = 2.05 + j * 0.9
                panel(slide, 0.73, y + 0.13, 0.1, 0.1, BLUE)
                box(slide, value, 1.05, y, 10.7, 0.77)
        elif layout == "comparison":
            for side, x in (("left", 0.7), ("right", 6.8)):
                panel(slide, x, 2.05, 5.8, 4.6)
                box(slide, body[f"{side}_title"], x + 0.38, 2.35, 5.02, 0.83,
                    size=22, color=BLUE, bold=True)
                for j, value in enumerate(body[f"{side}_items"]):
                    y = 3.35 + j * 0.73
                    panel(slide, x + 0.38, y + 0.09, 0.08, 0.08, BLUE)
                    box(slide, value, x + 0.7, y, 4.62, 0.64, size=18)
        elif layout == "metrics":
            items = body["items"]
            w = (11.9 - 0.3 * (len(items) - 1)) / len(items)
            for j, item in enumerate(items):
                x = 0.7 + j * (w + 0.3)
                panel(slide, x, 2.05, w, 4.6)
                box(slide, item["value"], x + 0.3, 2.6, w - 0.6, 0.7, size=36,
                    color=BLUE, bold=True, lines=1)
                box(slide, item["label"], x + 0.3, 3.7, w - 0.6, 0.85, size=22, bold=True)
                box(slide, item["context"], x + 0.3, 5.0, w - 0.6, 1.0,
                    size=18, color=MUTED, lines=3)
        elif layout == "table":
            rows, cols = body["rows"], body["columns"]
            table = slide.shapes.add_table(
                len(rows) + 1, len(cols), Inches(0.7), Inches(2.05), Inches(11.9),
                Inches(0.64 * (len(rows) + 1)),
            ).table
            for row_index, values in enumerate([cols] + rows):
                table.rows[row_index].height = Inches(0.64)
                for col_index, value in enumerate(values):
                    cell = table.cell(row_index, col_index)
                    cell.margin_left = cell.margin_right = Inches(0.12)
                    cell.margin_top = cell.margin_bottom = Inches(0.06)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(
                        BLUE if row_index == 0 else PANEL if row_index % 2 else WHITE
                    )
                    value = wrapped(value, 11.9 / len(cols) - 0.24, 14, 2, "table cell") if value else ""
                    write_frame(cell.text_frame, value, 14, WHITE if row_index == 0 else DARK, row_index == 0)
        elif layout == "chart":
            box(slide, f"Unit / Единица: {body['unit']}", 0.7, 1.92, 11.9, 0.38,
                size=18, color=MUTED, lines=1)
            chart_data = CategoryChartData()
            chart_data.categories = body["categories"]
            for series in body["series"]:
                chart_data.add_series(series["name"], series["values"])
            chart = slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(2.42),
                Inches(11.9), Inches(4.22), chart_data,
            ).chart
            chart.part.chart_workbook.update_from_xlsx_blob(literal_chart_workbook(body))
            chart.has_title = False
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            font_style(chart.legend.font, 18)
            font_style(chart.font, 18)
            category, axis = chart.category_axis, chart.value_axis
            category.tick_label_position = XL_TICK_LABEL_POSITION.LOW
            category.major_tick_mark = category.minor_tick_mark = XL_TICK_MARK.NONE
            axis.major_tick_mark = axis.minor_tick_mark = XL_TICK_MARK.NONE
            font_style(category.tick_labels.font, 18)
            font_style(axis.tick_labels.font, 18)
            category.tick_labels.offset = 100
            for element in category._element.xpath("./c:tickLblSkip | ./c:noMultiLvlLbl"):
                category._element.remove(element)
            single_level = OxmlElement("c:noMultiLvlLbl")
            single_level.set("val", "1")
            category._element.insert_element_before(single_level, "c:extLst")
            skip = OxmlElement("c:tickLblSkip")
            skip.set("val", "1")
            category._element.insert_element_before(skip, "c:noMultiLvlLbl", "c:extLst")
            values = [value for series in body["series"] for value in series["values"]]
            low, high = min(0, min(values)), max(0, max(values))
            span = high - low or 1
            axis.minimum_scale = low - span * 0.12 if low < 0 else 0
            axis.maximum_scale = high + span * 0.12 if high > 0 else (0 if low < 0 else 1)
            largest = max(abs(low), abs(high))
            axis.tick_labels.number_format = "0.##E+00" if largest >= 10000 or 0 < largest < 0.01 else "0.##"
            axis.tick_labels.number_format_is_linked = False
            axis.has_major_gridlines = True
            axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(PANEL)
            for j, series in enumerate(chart.series):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = RGBColor.from_string((BLUE, DARK, MUTED)[j])
                series.format.line.fill.background()
            # Values remain in the embedded workbook; dense data labels can collide.
            chart.plots[0].has_data_labels = False
            chart.plots[0].gap_width = 70
        elif layout == "closing":
            panel(slide, 0.7, 2.05, 11.9, 1.9)
            box(slide, body["takeaway"], 1.2, 2.3, 10.7, 1.4, size=26, bold=True, lines=3)
            for j, value in enumerate(body["actions"]):
                y = 4.3 + j * 0.77
                box(slide, f"{j + 1:02}", 0.7, y, 0.5, 0.4, size=20, color=BLUE, lines=1)
                box(slide, value, 1.3, y, 10.5, 0.67, size=20)
        box(slide, source_footer(data["source_ids"]), 0.7, 6.97, 10.7, 0.24,
            size=12, color=MUTED, lines=1)
        box(slide, f"{i + 1:02} / {len(content['slides']):02}", 11.55, 6.97, 1.05, 0.24,
            size=12, color=MUTED, lines=1).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        slide.notes_slide.notes_text_frame.text = notes_text(data, planned, outline, sources, placeholder)
    result = BytesIO()
    deck.save(result)
    return result.getvalue()


def markdown_escape(value: str) -> str:
    return re.sub(r"([\\`*_{}\[\]<>()#!|])", r"\\\1", value)


def markdown_documents(outline: dict, content: dict, placeholder=False) -> tuple[str, str]:
    source_index = {source["id"]: source for source in content["sources"]}
    title = markdown_escape(outline["title"])
    metadata = (
        f"# {title}\n\n"
        + (f"> {PLACEHOLDER}\n\n" if placeholder else "")
        + f"Language / Язык: {markdown_escape(outline['language'])}  \n"
        f"Audience / Аудитория: {markdown_escape(outline['audience'])}  \n"
        f"Mode: {outline['mode']}; outline version: {outline['version']}  \n"
        f"Planned duration: {sum(s['seconds'] for s in outline['slides']):g} seconds "
        f"(budget: {outline['duration_minutes']:g} minutes)\n\n"
    )
    if outline.get("theme") == "euclidean":
        metadata += "Design credit (not a research source):\n\n" + euclidean().DESIGN_CREDIT + "\n\n"
    speech = metadata
    for data, planned in zip(content["slides"], outline["slides"]):
        words = len((data["speech"] + " " + data["transition"]).split())
        speech += (
            f"## {data['id']} — {markdown_escape(data['title'])}\n\n"
            f"Timing / Время: {planned['seconds']:g} seconds. "
            f"Approximate whitespace-delimited word count: {words} "
            "(not a speech-speed validation).\n\n"
            f"### Speech / Речь\n\n{markdown_escape(data['speech'])}\n\n"
            f"### Transition / Переход\n\n{markdown_escape(data['transition']) or '—'}\n\n"
            "### Sources / Источники\n\n"
        )
        speech += "\n\n".join(markdown_escape(source_detail(source_index[ref])) for ref in data["source_ids"]) or "—"
        speech += "\n\n"
    sources = metadata + "## Sources / Источники\n\n"
    if not content["sources"]:
        sources += "No sources supplied / Источники не указаны.\n"
    for source in content["sources"]:
        used = ", ".join(slide["id"] for slide in content["slides"] if source["id"] in slide["source_ids"])
        sources += f"### {source['id']}\n\n{markdown_escape(source_detail(source))}\n\nUsed on slides: {used or '—'}\n\n"
    return speech, sources


def write_outputs(output_dir: Path, payloads: dict[str, bytes], overwrite=False) -> list[Path]:
    paths = [output_dir / name for name in payloads]
    for path in paths:
        if path.is_symlink() or (path.exists() and not path.is_file()):
            fail(str(path), "output must be a regular file, not a symlink or directory")
        if path.exists() and not overwrite:
            fail(str(path), "already exists; use --overwrite to replace outputs")
    output_dir.mkdir(parents=True, exist_ok=True)
    # Serialization finishes in memory before any outputs are opened. Exclusive
    # creation also protects against another process creating a file after preflight.
    created = []
    try:
        for path, payload in zip(paths, payloads.values()):
            with path.open("wb" if overwrite else "xb") as stream:
                created.append(path)
                stream.write(payload)
    except OSError:
        if not overwrite:
            for path in created:
                path.unlink()
        raise
    return paths


def render(outline: dict, content: dict, output_dir: Path, overwrite=False, placeholder=False) -> list[Path]:
    validate_outline(outline)
    validate_content(content, outline)
    pptx = build_presentation(outline, content, placeholder)
    speech, sources = markdown_documents(outline, content, placeholder)
    names = ("template.pptx", "template-speech.md", "template-sources.md") if placeholder else (
        "presentation.pptx", "speech.md", "sources.md"
    )
    return write_outputs(output_dir, dict(zip(names, (pptx, speech.encode("utf-8"), sources.encode("utf-8")))), overwrite)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    template = commands.add_parser("template", help="Render a clearly unapproved placeholder example")
    template.add_argument("--output-dir", type=Path, required=True)
    template.add_argument("--overwrite", action="store_true")
    template.add_argument("--theme", choices=("euclidean", "light-business"), default="euclidean")
    catalog = commands.add_parser("catalog", help="Print exact layout slots, styles, and geometry as JSON")
    catalog.add_argument("--theme", choices=("euclidean", "light-business"), default="euclidean")
    hashing = commands.add_parser("outline-hash", help="SHA256 of the exact outline file bytes")
    hashing.add_argument("--outline", type=Path, required=True)
    rendering = commands.add_parser("render", help="Render content bound to an explicitly approved outline")
    for key in ("outline", "approval", "content", "output-dir"):
        rendering.add_argument(f"--{key}", type=Path, required=True)
    rendering.add_argument("--overwrite", action="store_true")
    args = parser.parse_args(argv)
    try:
        if args.command == "outline-hash":
            print(hashlib.sha256(args.outline.read_bytes()).hexdigest())
            return 0
        if args.command == "catalog":
            data = euclidean().catalog_document() if args.theme == "euclidean" else {"theme": "light-business", "layouts": sorted(LAYOUTS)}
            print(json.dumps(data, ensure_ascii=False, indent=2))
            return 0
        if args.command == "template":
            assets = ROOT / "assets" / ("euclidean" if args.theme == "euclidean" else "")
            outline, _ = load_json(assets / "template-outline.json")
            content, _ = load_json(assets / "template-deck.json")
            paths = render(outline, content, args.output_dir, args.overwrite, placeholder=True)
        else:
            outline, raw = load_json(args.outline)
            validate_outline(outline)
            approval, _ = load_json(args.approval)
            validate_approval(approval, outline, raw)
            content, _ = load_json(args.content)
            paths = render(outline, content, args.output_dir, args.overwrite)
        for path in paths:
            print(path)
        return 0
    except (ValidationError, OSError, UnicodeError, json.JSONDecodeError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
